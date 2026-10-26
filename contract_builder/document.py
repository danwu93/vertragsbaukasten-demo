"""Bearbeitbares Word-Dokument ausschließlich im Arbeitsspeicher erzeugen."""
from io import BytesIO
import re
from docx import Document
from docx.shared import Cm, Pt
from .clauses import DISCLAIMER, sections


def generate_docx(contract):
    content = sections(contract)
    document = Document()
    document.core_properties.title = "Arbeitsvertrag – Muster"
    document.core_properties.author = "Vertragsbaukasten – fiktive Demo"
    page = document.sections[0]
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2.3)
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    document.add_heading("Arbeitsvertrag – Muster", 0)
    document.add_paragraph(DISCLAIMER, style="Subtitle")
    for title, text in content:
        document.add_heading(title, 1)
        document.add_paragraph(text)
    document.add_heading("Unterschriftenbereich – nur Demonstration", 1)
    document.add_paragraph("Fiktiver Ort / Datum: __________________________")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "____________________________"
    table.cell(0, 1).text = "____________________________"
    table.cell(1, 0).text = contract.company + " (Demo)"
    table.cell(1, 1).text = f"{contract.first_name} {contract.last_name} (Demo)"
    page.footer.paragraphs[0].text = "Vertragsbaukasten · Fiktiver Demonstrationsentwurf · Nicht rechtlich geprüft"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def filename(contract):
    name = re.sub(r"[^\w-]+", "_", f"{contract.first_name}_{contract.last_name}").strip("_")[:100]
    return f"Mustervertrag_{name or 'Demo'}.docx"
